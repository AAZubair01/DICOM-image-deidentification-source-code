#!/usr/bin/env python3
"""
dicom-cxr-deidentify - post-deidentification verification.

Run AFTER deidentify_dataset.py, on the same machine:
    python src/verify_deidentification.py --config config.yaml

For every DICOM under the output folder it walks ALL header elements -
including inside sequences, with full paths like
    RequestAttributesSequence[0].AccessionNumber
and reports the EXACT element holding any real name/code token from the
metadata spreadsheet (token length >= 4).

Philosophy: attack your own data before anyone else can. Any hit means
residual identifying information is present - do not share until it is
understood (most remaining hits after a correct run turn out to be
deliberately-retained institution names or coincidental digit collisions
inside hashed UIDs / binary blobs; the value snippet tells you which).

Options:
  --check-dates   also hunt for original exam dates (off by default; dates
                  are kept by policy and PS3.15 treats them as optional-shift)
  --ignore "A,B"  comma-separated tokens to exclude, e.g. institution words
                  you deliberately keep: --ignore "AMIN,KANO,TEACHING,HOSPITAL"

A clean run prints "0 potential identifier leaks".

Requirements:  pip install -r requirements.txt
"""

import os
import re
import csv
import sys
import argparse

try:
    import yaml
    import pydicom
    from openpyxl import load_workbook
except ImportError as exc:
    sys.exit("Missing dependency: %s\n"
             "Install with:  pip install -r requirements.txt" % exc)


def load_config(path):
    if not os.path.isfile(path):
        sys.exit("Config file not found: %s" % path)
    with open(path, encoding="utf-8") as fh:
        return yaml.safe_load(fh)


def norm(s):
    s = str(s).upper()
    s = re.sub(r"[^A-Z0-9 ]", " ", s)
    return " ".join(s.split())


def norm_compact(s):
    return norm(s).replace(" ", "")


def iter_elements(ds, path=""):
    """Yield (path, element) for every element, recursing into sequences.
    Never raises: unreadable tags/values are skipped."""
    try:
        elems = list(ds.elements())
    except Exception:
        return
    for elem in elems:
        try:
            kw = getattr(elem, "keyword", None) or str(getattr(elem, "tag", "?"))
        except Exception:
            kw = "(unreadable-tag)"
        p = "%s.%s" % (path, kw) if path else kw
        yield p, elem
        try:
            if getattr(elem, "VR", "") == "SQ":
                for i, item in enumerate(elem.value or []):
                    yield from iter_elements(item, "%s[%d]" % (p, i))
        except Exception:
            continue


def main(cfg, check_dates, ignore_tokens):
    out_dir = cfg["output"]["base_dir"]
    metadata_xlsx = cfg["project"]["metadata_file"]
    report_csv = os.path.join(out_dir, "verification_report.csv")

    print("Loading identifiers from metadata (all sheets) ...")
    print("  (name/code tokens always checked; dates: %s)"
          % ("checked" if check_dates else "NOT checked (use --check-dates)"))
    wb = load_workbook(metadata_xlsx, read_only=True, data_only=True)
    name_tokens, dates = set(), set()
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = [norm(h) if h is not None else "" for h in rows[0]]
        i_name_cols = [i for i, h in enumerate(header)
                       if ("NAME" in h or "CODE" in h) and "DATE" not in h]
        i_date = next((i for i, h in enumerate(header) if "DATE" in h), None)
        if not i_name_cols:
            continue
        for r in rows[1:]:
            if not r:
                continue
            ident = " ".join(str(r[i]) for i in i_name_cols
                             if i < len(r) and r[i] is not None)
            if not ident.strip() or norm(ident) in ("NAME CODE", "S N NAME CODE"):
                continue
            for tok in norm(ident).split():
                if len(tok) >= 4:      # short tokens like 'ABO' cause false alarms
                    name_tokens.add(tok)
            if check_dates and i_date is not None and i_date < len(r) \
                    and r[i_date] is not None:
                v = r[i_date]
                if hasattr(v, "strftime"):
                    dates.add(v.strftime("%Y%m%d"))
                else:
                    digits = re.sub(r"\D", "", str(v))
                    if len(digits) == 8:
                        dates.add(digits)
    wb.close()
    name_tokens -= set(norm(t) for t in ignore_tokens if t and t.strip())
    print("  %d name tokens%s to hunt for."
          % (len(name_tokens), (", %d dates" % len(dates)) if check_dates else ""))

    dcm_files = []
    for root, _, fns in os.walk(out_dir):
        for fn in fns:
            if fn.lower().endswith(".dcm"):
                dcm_files.append(os.path.join(root, fn))
    print("Scanning %d output DICOM files ..." % len(dcm_files))

    hits = []
    for n_file, fpath in enumerate(dcm_files):
        if (n_file + 1) % 25 == 0 or n_file == 0:
            print("  ... %d/%d" % (n_file + 1, len(dcm_files)))
        try:
            ds = pydicom.dcmread(fpath, stop_before_pixels=True)
        except Exception as exc:
            hits.append({"file": fpath, "tag": "-", "leaked": "-",
                         "value_snippet": "UNREADABLE FILE: %s" % exc})
            continue
        file_errors = 0
        for path, elem in iter_elements(ds):
            try:
                val = str(elem.value)
            except Exception:
                file_errors += 1
                continue
            nv, nc = norm(val), norm_compact(val)
            matched = None
            for tok in name_tokens:
                if tok in nv or tok in nc:
                    matched = tok
                    break
            if matched is None and check_dates:
                for d in dates:
                    if d and d in nv:
                        matched = d
                        break
            if matched is not None:
                hits.append({"file": fpath, "tag": path, "leaked": matched,
                             "value_snippet": val[:80]})
        if file_errors:
            hits.append({"file": fpath, "tag": "-", "leaked": "-",
                         "value_snippet": "%d element(s) unreadable (skipped)"
                         % file_errors})

    with open(report_csv, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=["file", "tag", "leaked",
                                           "value_snippet"])
        w.writeheader()
        w.writerows(hits)

    print("\n%d potential identifier leak(s)." % len(hits))
    if hits:
        by_tag = {}
        for h in hits:
            by_tag.setdefault(h["tag"], 0)
            by_tag[h["tag"]] += 1
        print("Hits by tag:")
        for tag, n in sorted(by_tag.items(), key=lambda kv: -kv[1]):
            print("    %-55s %d" % (tag, n))
        print("Details (with value snippets) written to:\n  %s" % report_csv)
        print("Do NOT share the dataset until every hit is understood and fixed.")
    else:
        print("Headers are clean of name/code tokens.")
        if not check_dates:
            print("(Original exam dates were kept by policy - run with "
                  "--check-dates to audit them too.)")
        print("Remaining risk is burned-in PIXEL text - visually inspect a")
        print("sample of images (and the qa_manual_review folder) before sharing.")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(
        description="Verify de-identified DICOM headers.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__)
    ap.add_argument("--config", default="config.yaml",
                    help="path to your config file (default: config.yaml)")
    ap.add_argument("--check-dates", action="store_true",
                    help="also hunt for original exam dates in headers")
    ap.add_argument("--ignore", default="",
                    help="comma-separated tokens to exclude from the hunt "
                         "(e.g. institution words you deliberately keep)")
    args = ap.parse_args()
    main(load_config(args.config), check_dates=args.check_dates,
         ignore_tokens=args.ignore.split(","))
